#!/usr/bin/env python3
import asyncio
import csv
import json
import math
import os
import random
import shutil
import subprocess
import sys
import wave
from pathlib import Path

import cv2
import edge_tts
import numpy as np
from docx import Document
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / ".claude" / "skills" / "jianying-editor"
OUT_DIR = ROOT / "workspace" / "jianying_corinthians_rosario_260820"
TTS_DIR = OUT_DIR / "voice_segments"
SCENE_DIR = OUT_DIR / "scene_stills"

DOCX_PATH = Path("/Users/jaguar/Downloads/lyz/Corinthians_vs_Rosario_Central_260820_逐字稿.docx")
LOGO_PATH = Path("/Users/jaguar/Downloads/lyz/20260820-084548.png")
ADBAR_PATH = Path("/Users/jaguar/Downloads/lyz/20260820-084541.jpg")
POSTER_PATH = Path("/Users/jaguar/Downloads/lyz/Codex 图像 2026年8月20日 11_39_47.png")
REFERENCE_PATH = Path(
    "/var/folders/xj/yk2515x93hv84d7rltqw9ckh0000gp/T/"
    "codex-clipboard-3158b80e-0bdb-4dfe-9c8c-4fc7db4fa9af.png"
)

PROJECT_NAME = "Corinthians_vs_Rosario_Central_260820"
WIDTH = 1080
HEIGHT = 1920
FPS = 30
VOICE = "pt-BR-FranciscaNeural"
RATE = "+4%"
PITCH = "+1Hz"

FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_IMPACT = "/System/Library/Fonts/Supplemental/Impact.ttf"


SCENE_PLANS = [
    "Poster de abertura com push-in, legenda sincronizada e transicao suave para a analise.",
    "Placar agregado 0-0 com alerta de decisao: quem piscar cai fora.",
    "Cartao de local e horario: 21h30, Neo Quimica Arena, Sao Paulo.",
    "Grafico limpo de volume: posse, finalizacoes e cruzamentos do Rosario.",
    "Alerta visual de jogador a menos e empate sustentado pelo Corinthians.",
    "Casa do Timao e bola parada como arma tática.",
    "Cartoes de ausencia: Allan suspenso, Andre fora, Yuri Alberto fora.",
    "Suspense Memphis: retorno gradual e impacto no segundo tempo.",
    "Ataque pelas pontas: Di Maria e Campaz abrindo corredores.",
    "Leitura do jogo: vantagem do Corinthians, mas sem conforto.",
    "Modelo probabilistico animado: 43%, 31%, 26%.",
    "Predicao de placar: Corinthians 1-0.",
    "Aviso forte: sem aposta recomendada.",
    "Medidor de risco alto para a odd do Rosario.",
    "Pergunta de interacao: decide em 90 minutos ou vai aos penaltis?",
    "Chamada natural para assistir com a galera no JaguarTV.",
    "Fechamento responsavel: analise ajuda, futebol nao promete nada.",
]


def run(cmd, *, check=True):
    print("+", " ".join(str(c) for c in cmd))
    return subprocess.run(cmd, check=check, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)


def ffprobe_duration(path: Path) -> float:
    result = run(
        [
            "ffprobe",
            "-v",
            "error",
            "-show_entries",
            "format=duration",
            "-of",
            "default=noprint_wrappers=1:nokey=1",
            str(path),
        ]
    )
    return float(result.stdout.strip())


def validate_inputs():
    required = {
        "transcript_docx": DOCX_PATH,
        "jaguartv_logo": LOGO_PATH,
        "bottom_ad_bar": ADBAR_PATH,
        "opening_poster": POSTER_PATH,
        "brand_reference": REFERENCE_PATH,
        "jianying_skill": SKILL_ROOT / "scripts" / "jy_wrapper.py",
    }
    missing = [f"{name}: {path}" for name, path in required.items() if not path.exists()]
    if missing:
        raise SystemExit("Missing required inputs:\n" + "\n".join(missing))

    for image_path in [LOGO_PATH, ADBAR_PATH, POSTER_PATH, REFERENCE_PATH]:
        with Image.open(image_path) as im:
            im.verify()

    Document(DOCX_PATH)


def extract_portuguese_sentences() -> list[str]:
    doc = Document(DOCX_PATH)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    try:
        start = paragraphs.index("9. 葡语—中文逐句对照全文") + 1
    except ValueError as exc:
        raise SystemExit("Could not find section 9 in DOCX.") from exc

    end = next((i for i in range(start, len(paragraphs)) if paragraphs[i].startswith("10.")), len(paragraphs))
    section = paragraphs[start:end]
    portuguese = section[0::2]
    chinese = section[1::2]
    if len(portuguese) != 17:
        raise SystemExit(f"Expected 17 Portuguese sentences, got {len(portuguese)}")
    if any(any("\u4e00" <= ch <= "\u9fff" for ch in line) for line in portuguese):
        raise SystemExit("Chinese text detected in extracted Portuguese list.")
    if len(chinese) != 17:
        raise SystemExit(f"Expected 17 Chinese check lines, got {len(chinese)}")
    return portuguese


def ensure_dirs():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TTS_DIR.mkdir(parents=True, exist_ok=True)
    SCENE_DIR.mkdir(parents=True, exist_ok=True)


async def generate_tts(sentences: list[str]):
    for idx, text in enumerate(sentences, 1):
        out = TTS_DIR / f"{idx:02d}.mp3"
        if out.exists() and out.stat().st_size > 1000:
            continue
        communicate = edge_tts.Communicate(text, VOICE, rate=RATE, pitch=PITCH)
        await communicate.save(str(out))


def convert_tts_to_wav(sentences: list[str]):
    segment_rows = []
    initial_silence = 0.15
    timeline_cursor = initial_silence
    pauses = []
    wav_paths = []
    for idx, sentence in enumerate(sentences, 1):
        mp3 = TTS_DIR / f"{idx:02d}.mp3"
        wav_path = TTS_DIR / f"{idx:02d}.wav"
        run(
            [
                "ffmpeg",
                "-y",
                "-i",
                str(mp3),
                "-ar",
                "48000",
                "-ac",
                "2",
                "-sample_fmt",
                "s16",
                str(wav_path),
            ]
        )
        dur = ffprobe_duration(wav_path)
        pause = 0.12 + (0.06 if sentence.endswith(".") else 0.0)
        if "," in sentence or ":" in sentence:
            pause += 0.04
        if idx in {7, 11, 13, 15}:
            pause += 0.06
        pause = min(0.30, max(0.12, pause))
        start = timeline_cursor
        end = start + dur
        segment_rows.append(
            {
                "index": idx,
                "text": sentence,
                "audio_file": str(mp3),
                "wav_file": str(wav_path),
                "duration": dur,
                "start": start,
                "end": end,
                "pause_after": pause,
                "scene": SCENE_PLANS[idx - 1],
            }
        )
        timeline_cursor = end + pause
        pauses.append(pause)
        wav_paths.append(wav_path)
    return segment_rows, wav_paths, pauses, timeline_cursor + 1.55


def read_wav_int16(path: Path):
    with wave.open(str(path), "rb") as wf:
        channels = wf.getnchannels()
        rate = wf.getframerate()
        frames = wf.readframes(wf.getnframes())
    audio = np.frombuffer(frames, dtype=np.int16).reshape(-1, channels)
    return audio, rate, channels


def write_wav_int16(path: Path, audio: np.ndarray, rate=48000):
    audio = np.clip(audio, -32768, 32767).astype(np.int16)
    with wave.open(str(path), "wb") as wf:
        wf.setnchannels(2)
        wf.setsampwidth(2)
        wf.setframerate(rate)
        wf.writeframes(audio.tobytes())


def assemble_narration(wav_paths: list[Path], pauses: list[float], total_duration: float):
    chunks = [np.zeros((int(48000 * 0.15), 2), dtype=np.int16)]
    for wav_path, pause in zip(wav_paths, pauses):
        audio, rate, channels = read_wav_int16(wav_path)
        if rate != 48000 or channels != 2:
            raise SystemExit(f"Unexpected WAV format: {wav_path}")
        chunks.append(audio)
        chunks.append(np.zeros((int(48000 * pause), 2), dtype=np.int16))
    chunks.append(np.zeros((int(48000 * 1.55), 2), dtype=np.int16))
    narration = np.vstack(chunks)
    expected = int(total_duration * 48000)
    if len(narration) < expected:
        narration = np.vstack([narration, np.zeros((expected - len(narration), 2), dtype=np.int16)])
    narration_path = OUT_DIR / f"{PROJECT_NAME}_narration.wav"
    write_wav_int16(narration_path, narration)
    return narration_path


def make_bgm(total_duration: float):
    sr = 48000
    t = np.linspace(0, total_duration, int(sr * total_duration), endpoint=False)
    bed = 0.035 * np.sin(2 * np.pi * 82 * t) + 0.025 * np.sin(2 * np.pi * 123 * t)
    pulse = np.zeros_like(t)
    beat_times = np.arange(0.4, total_duration, 0.62)
    for bt in beat_times:
        start = int(bt * sr)
        length = int(0.055 * sr)
        if start + length < len(pulse):
            env = np.exp(-np.linspace(0, 7, length))
            pulse[start : start + length] += 0.07 * env * np.sin(2 * np.pi * 148 * np.arange(length) / sr)
    shimmer = 0.008 * np.sin(2 * np.pi * 660 * t) * (0.5 + 0.5 * np.sin(2 * np.pi * 0.33 * t))
    audio = bed + pulse + shimmer
    audio[: int(sr * 1.0)] *= np.linspace(0, 1, int(sr * 1.0))
    audio[-int(sr * 1.2) :] *= np.linspace(1, 0, int(sr * 1.2))
    stereo = np.stack([audio, audio], axis=1)
    bgm_path = OUT_DIR / f"{PROJECT_NAME}_bgm_generated.wav"
    write_wav_int16(bgm_path, stereo * 32767)
    return bgm_path


def load_font(path, size):
    return ImageFont.truetype(path, size=size)


def text_size(draw, text, font):
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(draw, text, font, max_width, max_lines=2):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if text_size(draw, trial, font)[0] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    if len(lines) <= max_lines:
        return lines
    merged = lines[: max_lines - 1]
    tail = " ".join(lines[max_lines - 1 :])
    merged.append(tail)
    return merged


def draw_centered_text(draw, box, text, font, fill, stroke_width=0, stroke_fill=None, spacing=8):
    x0, y0, x1, y1 = box
    lines = wrap_text(draw, text, font, x1 - x0, 3)
    heights = [text_size(draw, line, font)[1] for line in lines]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = y0 + ((y1 - y0) - total_h) / 2
    for line, h in zip(lines, heights):
        w, _ = text_size(draw, line, font)
        draw.text(
            (x0 + (x1 - x0 - w) / 2, y),
            line,
            font=font,
            fill=fill,
            stroke_width=stroke_width,
            stroke_fill=stroke_fill,
        )
        y += h + spacing


def cover_image(path: Path, size=(WIDTH, HEIGHT), blur=False):
    im = Image.open(path).convert("RGB")
    scale = max(size[0] / im.width, size[1] / im.height)
    new_size = (int(im.width * scale), int(im.height * scale))
    im = im.resize(new_size, Image.Resampling.LANCZOS)
    left = (im.width - size[0]) // 2
    top = (im.height - size[1]) // 2
    im = im.crop((left, top, left + size[0], top + size[1]))
    if blur:
        im = im.filter(ImageFilter.GaussianBlur(12))
    return im


def contain_image(path: Path, max_w, max_h):
    im = Image.open(path).convert("RGBA")
    scale = min(max_w / im.width, max_h / im.height)
    return im.resize((int(im.width * scale), int(im.height * scale)), Image.Resampling.LANCZOS)


def base_background(frame_idx, total_frames, scene_idx):
    phase = frame_idx / FPS
    arr = np.zeros((HEIGHT, WIDTH, 3), dtype=np.uint8)
    y = np.linspace(0, 1, HEIGHT)[:, None]
    arr[:, :, 1] = (18 + 24 * (1 - y)).astype(np.uint8)
    arr[:, :, 0] = (3 + 10 * y).astype(np.uint8)
    arr[:, :, 2] = (8 + 18 * (1 - y)).astype(np.uint8)
    im = Image.fromarray(arr, "RGB")
    draw = ImageDraw.Draw(im, "RGBA")
    for i in range(80):
        rnd = random.Random(scene_idx * 1000 + i)
        x = int((rnd.random() * WIDTH + math.sin(phase * 0.5 + i) * 14) % WIDTH)
        y0 = int(60 + rnd.random() * 780)
        alpha = int(30 + rnd.random() * 65)
        draw.ellipse((x, y0, x + 2, y0 + 2), fill=(255, 235, 160, alpha))
    for x in range(-160, WIDTH + 160, 180):
        draw.line((x, 1040, x + 430, 1580), fill=(25, 110, 58, 72), width=4)
    draw.rectangle((0, 0, WIDTH, HEIGHT), outline=(0, 0, 0, 0))
    return im


def draw_panel(draw, xy, fill=(7, 17, 13, 190), outline=(246, 195, 62, 210), radius=22, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_scene(im, scene_idx, progress):
    draw = ImageDraw.Draw(im, "RGBA")
    title_font = load_font(FONT_IMPACT, 72)
    h1 = load_font(FONT_IMPACT, 110)
    h2 = load_font(FONT_BOLD, 56)
    small = load_font(FONT_BOLD, 36)
    regular = load_font(FONT_REGULAR, 34)
    gold = (246, 195, 62, 255)
    green = (62, 235, 127, 255)
    white = (245, 248, 242, 255)
    red = (238, 77, 66, 255)
    cyan = (94, 198, 255, 255)

    draw_panel(draw, (84, 92, 500, 152), fill=(6, 18, 14, 150), radius=28)
    draw.text((122, 106), "LIBERTADORES • ANÁLISE", font=small, fill=gold)

    if scene_idx == 1:
        draw_panel(draw, (130, 305, 950, 680))
        draw_centered_text(draw, (160, 328, 920, 470), "CORINTHIANS x ROSARIO CENTRAL", title_font, white)
        draw_centered_text(draw, (180, 480, 900, 645), "0–0  •  DECISÃO", h1, gold, 2, (0, 0, 0, 255))
    elif scene_idx == 2:
        draw_centered_text(draw, (100, 250, 980, 360), "HOJE • 21H30 BRT", title_font, gold)
        draw_panel(draw, (165, 395, 915, 650))
        draw_centered_text(draw, (210, 420, 870, 535), "NEO QUÍMICA ARENA", h2, white)
        draw_centered_text(draw, (210, 535, 870, 635), "SÃO PAULO", h1, gold)
    elif scene_idx == 3:
        labels = [("POSSE", 0.58), ("FINALIZAÇÕES", 0.72), ("CRUZAMENTOS", 0.84)]
        draw.text((110, 250), "ROSARIO COM MAIS VOLUME", font=title_font, fill=white)
        for i, (label, val) in enumerate(labels):
            y = 395 + i * 112
            draw.text((130, y), label, font=small, fill=white)
            draw.rounded_rectangle((130, y + 48, 895, y + 78), radius=15, fill=(255, 255, 255, 38))
            draw.rounded_rectangle((130, y + 48, 130 + int(765 * val * progress), y + 78), radius=15, fill=gold)
    elif scene_idx == 4:
        draw.text((120, 270), "UM A MENOS", font=h1, fill=red, stroke_width=2, stroke_fill=(0, 0, 0))
        draw_panel(draw, (180, 430, 900, 625), outline=(238, 77, 66, 220))
        draw_centered_text(draw, (210, 452, 870, 595), "CORINTHIANS SEGUROU O 0–0", h2, white)
        draw.rectangle((690, 245, 765, 350), fill=red)
    elif scene_idx == 5:
        draw.text((120, 245), "TIMÃO EM CASA", font=h1, fill=white, stroke_width=2, stroke_fill=(0, 0, 0))
        draw.arc((190, 430, 880, 840), start=200, end=340, fill=gold, width=7)
        draw.ellipse((522, 600, 558, 636), fill=white)
        draw.line((540, 618, 760, 510), fill=green, width=8)
        draw.text((145, 675), "BOLA PARADA PODE PESAR", font=h2, fill=gold)
    elif scene_idx == 6:
        draw.text((126, 230), "ALERTA DE DESFALQUES", font=title_font, fill=gold)
        names = [("ALLAN", "SUSPENSO"), ("ANDRÉ", "FORA"), ("YURI ALBERTO", "FORA")]
        for i, (name, status) in enumerate(names):
            y = 355 + i * 135
            draw_panel(draw, (150, y, 930, y + 105), outline=(238, 77, 66, 230))
            draw.text((190, y + 20), name, font=h2, fill=white)
            draw.text((705, y + 25), status, font=small, fill=red)
    elif scene_idx == 7:
        draw.text((135, 245), "MEMPHIS", font=h1, fill=gold, stroke_width=2, stroke_fill=(0, 0, 0))
        draw_centered_text(draw, (130, 405, 950, 585), "VOLTA AOS POUCOS?", h2, white)
        draw.rounded_rectangle((190, 625, 890, 660), radius=18, fill=(255, 255, 255, 45))
        draw.rounded_rectangle((190, 625, 190 + int(700 * min(1, progress)), 660), radius=18, fill=green)
        draw.text((270, 695), "SEGUNDO TEMPO", font=h2, fill=white)
    elif scene_idx == 8:
        draw.text((116, 240), "ATAQUE PELAS PONTAS", font=title_font, fill=white)
        draw.line((220, 630, 430, 430), fill=cyan, width=9)
        draw.line((860, 630, 650, 430), fill=gold, width=9)
        draw.ellipse((515, 500, 565, 550), fill=white)
        draw.text((130, 680), "DI MARÍA", font=h2, fill=cyan)
        draw.text((685, 680), "CAMPAZ", font=h2, fill=gold)
    elif scene_idx == 9:
        draw.text((112, 265), "LEITURA DO JOGO", font=title_font, fill=gold)
        draw_panel(draw, (130, 390, 950, 680))
        draw_centered_text(draw, (160, 420, 920, 535), "CORINTHIANS MAIS CHANCE", h2, white)
        draw_centered_text(draw, (160, 545, 920, 645), "MAS NÃO CONFORTÁVEL", h2, red)
    elif scene_idx == 10:
        probs = [("Corinthians", 43, gold), ("Empate", 31, white), ("Rosario", 26, cyan)]
        draw.text((130, 240), "MODELO DO DIA", font=title_font, fill=white)
        for i, (name, pct, color) in enumerate(probs):
            x = 92 + i * 330
            draw_panel(draw, (x, 380, x + 290, 690))
            draw_centered_text(draw, (x + 15, 405, x + 275, 470), name, small, white)
            draw_centered_text(draw, (x + 15, 485, x + 275, 600), f"{int(pct * progress)}%", h1, color)
            draw.rounded_rectangle((x + 45, 625, x + 245, 650), radius=12, fill=(255, 255, 255, 45))
            draw.rounded_rectangle((x + 45, 625, x + 45 + int(200 * pct / 50 * progress), 650), radius=12, fill=color)
    elif scene_idx == 11:
        draw.text((135, 245), "PLACAR PROVÁVEL", font=title_font, fill=white)
        draw_panel(draw, (165, 390, 915, 650))
        draw_centered_text(draw, (190, 410, 890, 620), "CORINTHIANS 1–0", h1, gold, 2, (0, 0, 0))
    elif scene_idx == 12:
        draw.text((145, 240), "SEM APOSTA", font=h1, fill=red, stroke_width=2, stroke_fill=(0, 0, 0))
        draw_centered_text(draw, (150, 420, 930, 590), "RECOMENDADA", h1, white, 2, (0, 0, 0))
        draw.line((170, 660, 910, 660), fill=red, width=9)
    elif scene_idx == 13:
        draw.text((130, 250), "RISCO ALTO", font=h1, fill=red, stroke_width=2, stroke_fill=(0, 0, 0))
        draw_panel(draw, (150, 440, 930, 585))
        draw.rounded_rectangle((210, 500, 870, 535), radius=18, fill=(255, 255, 255, 45))
        draw.rounded_rectangle((210, 500, 210 + int(610 * progress), 535), radius=18, fill=red)
        draw.text((230, 610), "ODD INTERESSANTE ≠ BOA ENTRADA", font=small, fill=white)
    elif scene_idx == 14:
        draw.text((135, 245), "COMENTA AÍ", font=h1, fill=gold, stroke_width=2, stroke_fill=(0, 0, 0))
        draw_panel(draw, (110, 430, 500, 620))
        draw_panel(draw, (580, 430, 970, 620))
        draw_centered_text(draw, (130, 450, 480, 595), "90 MINUTOS", h2, white)
        draw_centered_text(draw, (600, 450, 950, 595), "PÊNALTIS", h2, white)
    elif scene_idx == 15:
        draw.text((120, 245), "JUNTA A GALERA", font=title_font, fill=white)
        draw.text((145, 365), "E ASSISTE NO", font=h2, fill=gold)
        draw.text((145, 455), "JAGUARTV", font=h1, fill=green, stroke_width=2, stroke_fill=(0, 0, 0))
    elif scene_idx == 16:
        draw.text((120, 270), "ANÁLISE AJUDA", font=title_font, fill=gold)
        draw_centered_text(draw, (120, 420, 960, 610), "MAS FUTEBOL NÃO PROMETE NADA", h2, white)
    else:
        draw.text((120, 280), "CORINTHIANS x ROSARIO", font=title_font, fill=white)


def make_subtitle_blocks(rows):
    blocks = []
    block_id = 1
    for row in rows:
        text = row["text"]
        clauses = []
        current = ""
        for part in text.replace(": ", ":|").replace(", ", ",|").replace(". ", ".|").split("|"):
            part = part.strip()
            if not part:
                continue
            trial = part if not current else f"{current} {part}"
            if len(trial) <= 58:
                current = trial
            else:
                if current:
                    clauses.append(current)
                current = part
        if current:
            clauses.append(current)
        weights = [max(8, len(c)) for c in clauses]
        total_weight = sum(weights)
        cursor = row["start"]
        for i, (clause, weight) in enumerate(zip(clauses, weights)):
            if i == len(clauses) - 1:
                end = row["end"]
            else:
                end = cursor + row["duration"] * weight / total_weight
            blocks.append(
                {
                    "id": block_id,
                    "sentence_index": row["index"],
                    "start": cursor,
                    "end": end,
                    "text": clause,
                }
            )
            block_id += 1
            cursor = end
    return blocks


def srt_time(seconds):
    ms = int(round(seconds * 1000))
    h = ms // 3_600_000
    ms %= 3_600_000
    m = ms // 60_000
    ms %= 60_000
    s = ms // 1000
    ms %= 1000
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def write_srt(blocks):
    srt_path = OUT_DIR / f"{PROJECT_NAME}.srt"
    with srt_path.open("w", encoding="utf-8") as f:
        for block in blocks:
            f.write(f"{block['id']}\n")
            f.write(f"{srt_time(block['start'])} --> {srt_time(block['end'])}\n")
            f.write(block["text"] + "\n\n")
    return srt_path


def active_row(rows, t):
    for row in rows:
        if row["start"] <= t <= row["end"]:
            return row
    prior = [row for row in rows if row["start"] <= t]
    return prior[-1] if prior else rows[0]


def active_subtitle(blocks, t):
    for block in blocks:
        if block["start"] <= t <= block["end"]:
            return block
    return None


def paste_brand(im, t, total_duration):
    if t < 2.0:
        return
    logo = contain_image(LOGO_PATH, 116, 116)
    alpha = logo.split()[-1].point(lambda a: int(a * 0.88))
    logo.putalpha(alpha)
    im.alpha_composite(logo, (WIDTH - logo.width - 50, 62))
    bar = contain_image(ADBAR_PATH, WIDTH, int(HEIGHT * 0.072))
    y = HEIGHT - bar.height - 8
    fade_in = min(1.0, max(0.0, (t - 2.0) / 0.45))
    fade_out = min(1.0, max(0.0, (total_duration - t) / 0.75))
    a = min(fade_in, fade_out)
    bar_alpha = bar.split()[-1] if bar.mode == "RGBA" else Image.new("L", bar.size, 255)
    bar.putalpha(bar_alpha.point(lambda v: int(v * a)))
    im.alpha_composite(bar, ((WIDTH - bar.width) // 2, y))


def draw_subtitle(im, block):
    if not block:
        return
    draw = ImageDraw.Draw(im, "RGBA")
    font_size = 62 if len(block["text"]) < 76 else 55
    font = load_font(FONT_BOLD, font_size)
    max_width = int(WIDTH * 0.88)
    lines = wrap_text(draw, block["text"], font, max_width, 2)
    line_heights = [text_size(draw, line, font)[1] for line in lines]
    total_h = sum(line_heights) + 12 * (len(lines) - 1)
    center_y = int(HEIGHT * 0.745)
    y = center_y - total_h // 2
    box_pad_x = 28
    box_pad_y = 22
    widest = max(text_size(draw, line, font)[0] for line in lines)
    x0 = (WIDTH - widest) // 2 - box_pad_x
    y0 = y - box_pad_y
    x1 = (WIDTH + widest) // 2 + box_pad_x
    y1 = y + total_h + box_pad_y
    draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=(0, 0, 0, 168))
    for line, h in zip(lines, line_heights):
        w, _ = text_size(draw, line, font)
        x = (WIDTH - w) // 2
        draw.text((x, y), line, font=font, fill=(255, 255, 255, 255), stroke_width=4, stroke_fill=(0, 0, 0, 255))
        y += h + 12


def render_video(rows, blocks, total_duration, *, captions: bool):
    suffix = "noaudio_subtitled" if captions else "noaudio_clean"
    noaudio_path = OUT_DIR / f"{PROJECT_NAME}_{suffix}.mp4"
    writer = cv2.VideoWriter(str(noaudio_path), cv2.VideoWriter_fourcc(*"mp4v"), FPS, (WIDTH, HEIGHT))
    total_frames = int(math.ceil(total_duration * FPS))
    poster_base = cover_image(POSTER_PATH).convert("RGBA")
    for frame in range(total_frames):
        t = frame / FPS
        if t < 2.0:
            zoom = 1.0 + 0.025 * min(1, t / 2.0)
            w = int(WIDTH / zoom)
            h = int(HEIGHT / zoom)
            left = (WIDTH - w) // 2
            top = (HEIGHT - h) // 2
            im = poster_base.crop((left, top, left + w, top + h)).resize((WIDTH, HEIGHT), Image.Resampling.LANCZOS)
            if t < 0.20:
                alpha = t / 0.20
                black = Image.new("RGBA", (WIDTH, HEIGHT), (0, 0, 0, 255))
                im = Image.blend(black, im, alpha)
            if t > 1.60:
                blur = min(4.0, (t - 1.60) / 0.40 * 4.0)
                im = im.filter(ImageFilter.GaussianBlur(blur))
                overlay = Image.new("RGBA", (WIDTH, HEIGHT), (246, 195, 62, int((t - 1.60) / 0.40 * 35)))
                im = Image.alpha_composite(im, overlay)
        else:
            row = active_row(rows, t)
            scene_idx = max(1, row["index"] - 1)
            scene_start = max(2.0, row["start"])
            progress = min(1.0, max(0.0, (t - scene_start) / max(0.6, row["duration"])))
            im = base_background(frame, total_frames, scene_idx).convert("RGBA")
            draw_scene(im, scene_idx, progress)
            vignette = Image.new("L", (WIDTH, HEIGHT), 0)
            vd = ImageDraw.Draw(vignette)
            vd.ellipse((-260, -80, WIDTH + 260, HEIGHT + 100), fill=210)
            vignette = vignette.filter(ImageFilter.GaussianBlur(120))
            dark = Image.new("RGBA", (WIDTH, HEIGHT), (0, 0, 0, 120))
            im = Image.composite(im, Image.alpha_composite(im, dark), vignette)
            paste_brand(im, t, total_duration)
        if captions:
            draw_subtitle(im, active_subtitle(blocks, t))
        if frame % max(1, total_frames // 12) == 0:
            print(f"render {frame}/{total_frames}")
        writer.write(cv2.cvtColor(np.array(im.convert("RGB")), cv2.COLOR_RGB2BGR))
    writer.release()
    return noaudio_path


def mix_final_video(noaudio_path: Path, narration_path: Path, bgm_path: Path):
    final_path = OUT_DIR / f"{PROJECT_NAME}.mp4"
    run(
        [
            "ffmpeg",
            "-y",
            "-i",
            str(noaudio_path),
            "-i",
            str(narration_path),
            "-i",
            str(bgm_path),
            "-filter_complex",
            "[1:a]volume=1.0[nar];"
            "[2:a]volume=0.16[bgm];"
            "[nar][bgm]amix=inputs=2:duration=first:dropout_transition=0,"
            "loudnorm=I=-16:TP=-1:LRA=11,"
            "volume=0.89,"
            "alimiter=limit=0.891[a]",
            "-map",
            "0:v:0",
            "-map",
            "[a]",
            "-c:v",
            "libx264",
            "-preset",
            "medium",
            "-crf",
            "18",
            "-pix_fmt",
            "yuv420p",
            "-r",
            str(FPS),
            "-c:a",
            "aac",
            "-ar",
            "48000",
            "-b:a",
            "192k",
            "-movflags",
            "+faststart",
            "-shortest",
            str(final_path),
        ]
    )
    return final_path


def write_metadata(sentences, rows, blocks, total_duration):
    (OUT_DIR / "portuguese_sentences_17.txt").write_text(
        "\n".join(f"{i+1:02d}. {s}" for i, s in enumerate(sentences)) + "\n", encoding="utf-8"
    )
    with (OUT_DIR / "timeline_segments.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "index",
                "text",
                "audio_file",
                "duration_ms",
                "start_time",
                "end_time",
                "pause_after_ms",
                "subtitle_blocks",
                "scene",
            ],
        )
        writer.writeheader()
        for row in rows:
            row_blocks = [str(b["id"]) for b in blocks if b["sentence_index"] == row["index"]]
            writer.writerow(
                {
                    "index": row["index"],
                    "text": row["text"],
                    "audio_file": row["audio_file"],
                    "duration_ms": round(row["duration"] * 1000),
                    "start_time": srt_time(row["start"]),
                    "end_time": srt_time(row["end"]),
                    "pause_after_ms": round(row["pause_after"] * 1000),
                    "subtitle_blocks": ",".join(row_blocks),
                    "scene": row["scene"],
                }
            )
    (OUT_DIR / "timeline_segments.json").write_text(
        json.dumps({"total_duration": total_duration, "sentences": rows, "subtitle_blocks": blocks}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def create_scene_stills(rows, blocks, total_duration):
    times = [0.4, 2.3]
    for idx in [10, 11, 14, 16]:
        times.append(rows[idx]["start"] + 0.45)
    for n, t in enumerate(times, 1):
        if t < 2:
            im = cover_image(POSTER_PATH).convert("RGBA")
        else:
            row = active_row(rows, t)
            scene_idx = max(1, row["index"] - 1)
            im = base_background(int(t * FPS), int(total_duration * FPS), scene_idx).convert("RGBA")
            draw_scene(im, scene_idx, 1.0)
            paste_brand(im, t, total_duration)
        draw_subtitle(im, active_subtitle(blocks, t))
        out = SCENE_DIR / f"check_{n:02d}_{t:.2f}s.png"
        im.convert("RGB").save(out, quality=95)


def create_jianying_draft(noaudio_path: Path, narration_path: Path, bgm_path: Path, srt_path: Path, total_duration: float):
    current_dir = str(ROOT)
    env_root = os.getenv("JY_SKILL_ROOT", "").strip()
    skill_candidates = [
        env_root,
        os.path.join(current_dir, ".agent", "skills", "jianying-editor"),
        os.path.join(current_dir, ".trae", "skills", "jianying-editor"),
        os.path.join(current_dir, ".claude", "skills", "jianying-editor"),
        os.path.join(current_dir, "skills", "jianying-editor"),
        os.path.abspath(".agent/skills/jianying-editor"),
        os.path.abspath(".trae/skills/jianying-editor"),
        os.path.abspath(".claude/skills/jianying-editor"),
        os.path.abspath("skills/jianying-editor"),
    ]
    scripts_path = None
    attempted = []
    for p in skill_candidates:
        if not p:
            continue
        p = os.path.abspath(p)
        attempted.append(p)
        if os.path.exists(os.path.join(p, "scripts", "jy_wrapper.py")):
            scripts_path = os.path.join(p, "scripts")
            break
    if not scripts_path:
        raise ImportError("Could not find jianying-editor/scripts/jy_wrapper.py\nTried:\n- " + "\n- ".join(attempted))
    if scripts_path not in sys.path:
        sys.path.insert(0, scripts_path)
    from jy_wrapper import JyProject, draft

    project = JyProject(PROJECT_NAME, width=WIDTH, height=HEIGHT, overwrite=True)
    project.add_media_safe(str(noaudio_path.resolve()), start_time="0s", duration=f"{total_duration:.3f}s", track_name="RenderedVisual")
    project.add_audio_safe(str(narration_path.resolve()), start_time="0s", duration=f"{total_duration:.3f}s", track_name="Narration")
    bgm = project.add_audio_safe(str(bgm_path.resolve()), start_time="0s", duration=f"{total_duration:.3f}s", track_name="BGM")
    if bgm:
        bgm.volume = 0.16
    project.script.import_srt(
        str(srt_path.resolve()),
        "Portuguese_Subtitles",
        text_style=draft.TextStyle(size=5.8, align=1, auto_wrapping=True),
        clip_settings=draft.ClipSettings(transform_y=-0.48),
    )
    result = project.save()
    (OUT_DIR / "jianying_draft_result.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result


def validate_outputs(final_path: Path, noaudio_path: Path, narration_path: Path, srt_path: Path, total_duration: float):
    final_duration = ffprobe_duration(final_path)
    nar_duration = ffprobe_duration(narration_path)
    if abs(final_duration - nar_duration) > 0.25:
        raise SystemExit(f"Audio/video duration mismatch: final={final_duration:.3f}, narration={nar_duration:.3f}")
    if final_duration < 20:
        raise SystemExit("Final video unexpectedly short.")
    if final_path.stat().st_size < 1_000_000:
        raise SystemExit("Final video file is too small.")
    if not srt_path.exists() or srt_path.stat().st_size < 100:
        raise SystemExit("SRT missing or empty.")
    black_check = OUT_DIR / "blackdetect.txt"
    result = run(
        [
            "ffmpeg",
            "-i",
            str(final_path),
            "-vf",
            "blackdetect=d=0.3:pic_th=0.98",
            "-an",
            "-f",
            "null",
            "-",
        ],
        check=False,
    )
    black_check.write_text(result.stdout, encoding="utf-8")
    return {"final_duration": final_duration, "narration_duration": nar_duration, "target_duration": total_duration}


def main():
    random.seed(260820)
    ensure_dirs()
    validate_inputs()
    sentences = extract_portuguese_sentences()
    print("Portuguese sentences:")
    for idx, sentence in enumerate(sentences, 1):
        print(f"{idx:02d}. {sentence}")
    asyncio.run(generate_tts(sentences))
    rows, wav_paths, pauses, total_duration = convert_tts_to_wav(sentences)
    narration_path = assemble_narration(wav_paths, pauses, total_duration)
    bgm_path = make_bgm(total_duration)
    blocks = make_subtitle_blocks(rows)
    srt_path = write_srt(blocks)
    write_metadata(sentences, rows, blocks, total_duration)
    subtitled_noaudio_path = render_video(rows, blocks, total_duration, captions=True)
    clean_noaudio_path = render_video(rows, blocks, total_duration, captions=False)
    final_path = mix_final_video(subtitled_noaudio_path, narration_path, bgm_path)
    create_scene_stills(rows, blocks, total_duration)
    draft_result = create_jianying_draft(clean_noaudio_path, narration_path, bgm_path, srt_path, total_duration)
    validation = validate_outputs(final_path, clean_noaudio_path, narration_path, srt_path, total_duration)
    summary = {
        "final_video": str(final_path),
        "noaudio_subtitled_video": str(subtitled_noaudio_path),
        "noaudio_clean_video": str(clean_noaudio_path),
        "narration": str(narration_path),
        "bgm": str(bgm_path),
        "srt": str(srt_path),
        "timeline_csv": str(OUT_DIR / "timeline_segments.csv"),
        "timeline_json": str(OUT_DIR / "timeline_segments.json"),
        "voice_segments_dir": str(TTS_DIR),
        "scene_stills_dir": str(SCENE_DIR),
        "jianying_draft": draft_result,
        "validation": validation,
    }
    (OUT_DIR / "delivery_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
